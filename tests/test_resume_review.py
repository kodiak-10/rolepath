from io import BytesIO
import pytest
from docx import Document
from pypdf import PdfWriter
from rolepath import resume_review as r


def test_docx_resume_reads_paragraph_and_table():
    doc=Document();doc.add_paragraph('项目经历：设计用户反馈流程并整理访谈记录。')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='能力';table.cell(0,1).text='数据分析'
    out=BytesIO();doc.save(out)
    result=r.extract_resume('resume.docx',out.getvalue())
    assert '用户反馈' in result and '数据分析' in result


def test_utf8_text_and_unreadable_documents():
    assert '产品实习' in r.extract_resume('resume.txt','产品实习经历：整理用户反馈，撰写需求文档。'.encode())
    with pytest.raises(ValueError,match='UTF-8'):r.extract_resume('resume.txt',b'\xff'*30)
    with pytest.raises(ValueError,match='无法读取'):r.extract_resume('resume.docx',b'broken content'*4)
    out=BytesIO();w=PdfWriter();w.add_blank_page(width=100,height=100);w.write(out)
    with pytest.raises(ValueError,match='扫描版'):r.extract_resume('resume.pdf',out.getvalue())


def test_resume_rejects_large_files_and_inputs():
    with pytest.raises(ValueError,match='5MB'):r.extract_resume('resume.txt',b'x'*(r.MAX_FILE_BYTES+1))
    with pytest.raises(ValueError,match='20000'):r.extract_resume('resume.txt',b'x'*(r.MAX_TEXT_CHARS+1))
    with pytest.raises(ValueError):r.extract_resume('resume.exe',b'x'*30)


def test_review_sends_resume_and_job_as_data_with_no_invented_facts(monkeypatch):
    from rolepath import llm
    import json
    seen={}
    def fake(content,system,**kwargs):
        seen.update(content=json.loads(content),system=system,kwargs=kwargs)
        return '请补充真实项目结果。',None
    monkeypatch.setattr(llm,'_message',fake)
    answer=r.review_resume('负责整理用户访谈记录以及协助产品需求分析。','给出修改建议','需要用户研究和需求分析',[], 'zh')
    assert answer['answer']=='请补充真实项目结果。'
    assert seen['content']['job_description']=='需要用户研究和需求分析'
    assert '不得编造' in seen['system'] and seen['kwargs']['timeout']==90
